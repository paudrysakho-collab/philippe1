#!/usr/bin/env python3
"""Le tableur des prix du salon, pré-rempli avec les prix du tarif de septembre 2026.

L'agence (3 octobre 2026) : « remettre les prix ». Pour chaque vin dégusté rapproché d'une
ligne du tarif (`tarif` dans data/salon-prive-2026.json), les colonnes de paliers reçoivent
les prix de cette ligne, tels quels. Un vin absent du tarif reste vide. Deux colonnes de plus
(J et K) disent, ligne par ligne, d'où vient le prix et ce qui mérite un coup d'œil (millésime,
couleur ou cuvée différents entre le dossier de Mathéo et le tarif). L'import
(npm run importer-salon) ne lit que les colonnes A et F à I : J et K ne gênent pas.

Rien n'est écrit dans les données : c'est une proposition, à valider par l'agence, puis à
importer. Sortie : tableur/prix-salon-prive-2026-propose.xlsx.
"""
import json
import pathlib
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import tableur as T  # noqa: E402

SORTIE = T.RACINE / "tableur/prix-salon-prive-2026-propose.xlsx"
ORANGE = "F8DDB8"   # à regarder
ROUGE = "F4C7C7"    # pas de prix au tarif
GENERIQUE = "absent du tarif de septembre"


def euros(c):
    return f"{c / 100:.2f}".replace(".", ",") + " €"


def main():
    T.exporter_salon(SORTIE)
    salon = json.loads(T.SALON.read_text(encoding="utf-8"))
    fiches = T.charger_fiches()
    vins = {f"S{st['stand']:02d} V{i:02d}": v
            for st in salon["stands"] for i, v in enumerate(st["vins"], start=1)}

    wb = load_workbook(SORTIE)
    ws = wb["Prix du salon"]
    ws["A2"] = ("PROPOSITION : prix repris du tarif de septembre 2026, à valider. Colonne J : d'où vient "
                "le prix ; orange = à regarder, rouge = pas de prix au tarif. Ne pas toucher à la colonne A.")
    for c, (titre, larg) in enumerate((("Prix repris ?", 34), ("Ligne du tarif (pour comparer)", 60)), start=10):
        cel = ws.cell(T.LIGNE_TITRES, c, titre)
        cel.font = Font(name=T.POLICE, size=10, bold=True, color=T.ENCRE)
        cel.fill = PatternFill("solid", fgColor=T.CRAIE)
        ws.column_dimensions[cel.column_letter].width = larg

    n_prix = n_orange = n_rouge = 0
    for row in ws.iter_rows(min_row=T.LIGNE_TITRES + 1):
        ref = T.texte(row[0].value)
        if ref not in vins:
            continue
        v, r = vins[ref], row[0].row
        a_voir = [e for e in v.get("ecarts", []) if not e.startswith(GENERIQUE)]
        if v.get("tarif"):
            ligne = fiches[v["fiche"]]["tableaux"][v["tarif"]["tableau"]]["lignes"][v["tarif"]["ligne"]]
            prix = ligne["prix_centimes"]
            for pi, c in enumerate(prix):
                if c is not None:
                    ws.cell(r, 7 + pi, c / 100)
            n_prix += 1
            nom = " — ".join(x for x in (ligne.get("appellation"), ligne.get("cuvee")) if x)
            ws.cell(r, 11, f"{nom} · {ligne.get('couleur') or '—'} · {ligne.get('millesime') or '—'} · "
                           f"{ligne.get('contenance') or '—'} · " + " / ".join(euros(c) for c in prix if c is not None))
            etat = "Oui, prix du tarif" + (" — à regarder : " + " ; ".join(a_voir) if a_voir else "")
            fond = ORANGE if a_voir else None
            n_orange += bool(a_voir)
        else:
            etat = "Non : vin absent du tarif — à remplir, ou laisser vide"
            if a_voir:
                etat += " (" + " ; ".join(a_voir) + ")"
            fond = ROUGE
            n_rouge += 1
        cel = ws.cell(r, 10, etat)
        for c in (10, 11):
            ws.cell(r, c).font = Font(name=T.POLICE, size=9, color=T.ENCRE)
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        if fond:
            for c in range(2, 12):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=fond)
    wb.save(SORTIE)
    print(f"✓ {SORTIE.relative_to(T.RACINE)} — {n_prix} vins avec les prix du tarif "
          f"(dont {n_orange} à regarder), {n_rouge} sans prix au tarif")


if __name__ == "__main__":
    main()
